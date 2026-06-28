from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("analysis/xsc-wmf-acceptance-report.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tbl_cell_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    set_cell_margins(table)


def set_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("XSC WMF acceptance report")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)


def style_document(doc):
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    set_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("XSC 自写 WMF 预览验收报告")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("项目：J:\\latextomathtype    日期：2026-06-14")
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("1F3A5F")
    doc.add_paragraph()


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "结果"
    for cell in hdr:
        set_cell_shading(cell, "F2F4F7")
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value

    set_table_geometry(table, [2300, 7060])

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10.5)
    doc.add_paragraph()


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.add_run(text)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_title(doc)

    add_callout(
        doc,
        "结论：XSC 1..551 全量重建验收通过。生成侧 153956 个 WMF 对象全部完成尺寸比较，最大目标宽误差 0.3440%，最大目标高误差 0.3298%，低于 1% 阈值；生成侧非 WMF 数为 0。",
    )

    doc.add_heading("验收范围", level=1)
    doc.add_paragraph(
        "本报告记录 J:\\latextomathtype 当前自写 WMF 预览线路的验收结果。测试集来自 E:\\新加卷\\新建文件夹\\xsc资料，覆盖编号 1 到 551 的教师版 DOCX。"
    )
    add_bullet(doc, "预览线路要求：坚持自写 WMF，不使用位图、DIB 或其他预览回退路线。")
    add_bullet(doc, "验收方式：逐分片重建 DOCX，再对生成文档与源文档的公式对象数、WMF 类型和物理尺寸进行比较。")
    add_bullet(doc, "目标阈值：生成侧所有可比较 WMF 对象宽高误差均小于 1%。")

    doc.add_heading("核心指标", level=1)
    add_key_value_table(
        doc,
        [
            ("文件覆盖", "1..551，覆盖完整，无缺口"),
            ("分片数量", "22 个分片，1..25 已重新跑干净 run"),
            ("源侧 WMF 对象", "153956 个；其中 153938 个带可比较尺寸字段，已比较对象全部在 1% 内"),
            ("目标侧 WMF 对象", "153956 个；宽高尺寸全部完成比较"),
            ("目标宽误差", "最大 0.3440366972%"),
            ("目标高误差", "最大 0.3298436828%"),
            ("生成侧非 WMF", "0"),
            ("验收结论", "通过，目标侧对象数、WMF 类型和物理尺寸均满足要求"),
        ],
    )

    doc.add_heading("验证命令", level=1)
    doc.add_paragraph("关键复跑命令：")
    add_bullet(
        doc,
        "powershell.exe -NoProfile -ExecutionPolicy Bypass -File scripts\\run_xsc_unattended_acceptance.ps1 -Start 1 -End 25 -ChunkSize 25 -CommandTimeoutMinutes 45",
    )
    add_bullet(
        doc,
        ".mvn\\apache-maven-3.9.12\\bin\\mvn.cmd -q \"-Dtest=MathIRConverterTest,VectorWmfFormulaRendererTest\" test",
    )
    doc.add_paragraph(
        "全量聚合基于 analysis\\unattended-runs 下 22 个 summary.json，目标侧 badChunks 为空，pass=true。"
    )

    doc.add_heading("实现说明", level=1)
    add_bullet(doc, "渲染侧补强了自写 Vector WMF 对常见函数命令、max、backsim 等公式符号的覆盖。")
    add_bullet(doc, "Tokenizer、MathIR 转换和 MTEF 字符映射同步补齐，避免测试集中公式命令落入未知或降级路径。")
    add_bullet(doc, "验收脚本支持无人值守分片运行，便于 Codex 闪退后继续从磁盘结果收口。")

    doc.add_heading("注意事项", level=1)
    doc.add_paragraph(
        "当前验收的强证据集中在对象数、WMF 类型和物理尺寸。内容层面通过同一 LaTeX 请求重建和代表性单元测试覆盖进行间接确认；未做逐像素视觉 diff。若后续需要更严格内容验收，可在现有分片输出上追加 WMF 渲染快照或 Office 预览图的逐页差异检测。"
    )

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
